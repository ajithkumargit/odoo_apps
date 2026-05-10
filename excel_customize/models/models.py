import copy

from odoo import api, models, fields

import base64
import io
from datetime import datetime
from openpyxl import load_workbook, Workbook

class ExcelData(models.Model):
    _name = 'excel.data'

    name = fields.Char()
    file = fields.Binary()
    file_type = fields.Many2one("file.type")

    def customizeExcel(self, vals):
        result = {'status':'failure','message':'File Not Loaded','binary':''}
        uploaded_file = vals['uploadedFIle']
        file_content = uploaded_file.read()

        # Load existing workbook from uploaded file
        old_wb = load_workbook(io.BytesIO(file_content))
        old_ws = old_wb.active

        # Create new workbook
        new_wb = Workbook()
        new_ws = new_wb.active
        new_ws.title = "Processed Statement"

        # Copy first 2 rows (header rows)
        for row in old_ws.iter_rows(min_row=1, max_row=2):
            for cell in row:

                new_cell = new_ws[cell.coordinate]
                new_cell.value = cell.value

                # Copy styles
                if cell.has_style:
                    new_cell.font = copy(cell.font)
                    new_cell.fill = copy(cell.fill)
                    new_cell.border = copy(cell.border)
                    new_cell.alignment = copy(cell.alignment)
                    new_cell.number_format = cell.number_format

        # Copy column widths
        for col, dim in old_ws.column_dimensions.items():
            new_ws.column_dimensions[col].width = dim.width

        # --------------------------------
        # ADD YOUR CUSTOM PROCESSING HERE
        # --------------------------------

        # Example:
        # new_ws["A3"] = "Processed by Odoo"

        # Save to memory
        output = io.BytesIO()
        new_wb.save(output)
        output.seek(0)

        new_filename = f"processed_statement-{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}.xlsx"

        return result
class NameToCategory(models.Model):
    _name = 'name.to.category'

    name = fields.Char()
    statement_detail_ids = fields.Many2many("statement.detail")

class FileType(models.Model):
    _name = 'file.type'

    name = fields.Char()

class StatementDetail(models.Model):
    _name = 'statement.detail'

    name = fields.Char()